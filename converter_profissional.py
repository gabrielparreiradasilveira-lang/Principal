#!/usr/bin/env python3
"""
PDF to PPTX Converter - Versão Profissional
Converte PDF para slides formatados profissionalmente
"""

from PyPDF2 import PdfReader
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import os
import sys


class ConversorPDFProfissional:
    def __init__(self, pdf_path: str):
        self.pdf_path = pdf_path
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)

    def extrair_pdf(self):
        """Extrai conteúdo do PDF"""
        print(f"📄 Extraindo: {os.path.basename(self.pdf_path)}...")

        paginas = []
        pdf_reader = PdfReader(self.pdf_path)

        for num, pagina in enumerate(pdf_reader.pages, 1):
            texto = pagina.extract_text()
            if texto and texto.strip():
                linhas = [l.strip() for l in texto.split('\n') if l.strip()]
                paginas.append({
                    'numero': num,
                    'titulo': linhas[0][:80] if linhas else f"Página {num}",
                    'conteudo': '\n'.join(linhas[1:]) if len(linhas) > 1 else texto,
                })

        print(f"✅ {len(paginas)} páginas extraídas\n")
        return paginas

    def criar_slide_titulo(self, titulo: str, subtitulo: str = ""):
        """Cria slide de título com estilo profissional"""
        blank_slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_slide_layout)

        # Background azul profissional
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(31, 78, 121)

        # Título grande
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5), Inches(9), Inches(2)
        )
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        title_frame.text = titulo

        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(54)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(255, 255, 255)
        title_para.alignment = PP_ALIGN.CENTER

        # Subtítulo
        if subtitulo:
            sub_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(4.5), Inches(9), Inches(1.5)
            )
            sub_frame = sub_box.text_frame
            sub_frame.text = subtitulo
            sub_para = sub_frame.paragraphs[0]
            sub_para.font.size = Pt(18)
            sub_para.font.color.rgb = RGBColor(200, 200, 200)
            sub_para.alignment = PP_ALIGN.CENTER

    def criar_slide_conteudo(self, titulo: str, conteudo: str, numero: int = 0, total: int = 0):
        """Cria slide com conteúdo bem formatado"""
        blank_slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_slide_layout)

        # Background branco
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(255, 255, 255)

        # Barra azul no topo
        shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(31, 78, 121)
        shape.line.color.rgb = RGBColor(31, 78, 121)

        # Título na barra
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(8.5), Inches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = titulo
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(32)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(255, 255, 255)
        title_para.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Número da página
        if numero > 0:
            num_box = slide.shapes.add_textbox(Inches(8.8), Inches(0.15), Inches(1), Inches(0.6))
            num_frame = num_box.text_frame
            num_frame.text = f"{numero}"
            num_para = num_frame.paragraphs[0]
            num_para.font.size = Pt(16)
            num_para.font.color.rgb = RGBColor(200, 200, 200)
            num_para.alignment = PP_ALIGN.RIGHT
            num_para.vertical_anchor = MSO_ANCHOR.MIDDLE

        # Conteúdo
        content_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.2), Inches(8.4), Inches(5.8)
        )
        text_frame = content_box.text_frame
        text_frame.word_wrap = True

        # Processa conteúdo
        linhas_conteudo = conteudo.split('\n')

        for i, linha in enumerate(linhas_conteudo):
            if i > 0:
                text_frame.add_paragraph()

            p = text_frame.paragraphs[i]
            p.text = linha.strip()
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(40, 40, 40)
            p.space_after = Pt(8)
            p.level = 0

            # Destaca linhas que parecem ser títulos
            if len(linha.strip()) < 60 and linha.isupper():
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = RGBColor(31, 78, 121)
                p.space_before = Pt(10)

    def converter(self, output_path: str = None):
        """Executa conversão completa"""
        if output_path is None:
            output_path = self.pdf_path.replace('.pdf', '_PROFISSIONAL.pptx')

        print("=" * 60)
        print("  PDF → Slides Profissional")
        print("=" * 60 + "\n")

        paginas = self.extrair_pdf()
        if not paginas:
            return None

        # Slide de título
        nome_arquivo = os.path.basename(self.pdf_path).replace('.pdf', '')
        self.criar_slide_titulo(
            nome_arquivo,
            "Convertido automaticamente"
        )

        # Slides de conteúdo
        for pagina in paginas:
            self.criar_slide_conteudo(
                pagina['titulo'],
                pagina['conteudo'],
                pagina['numero'],
                len(paginas)
            )

        # Salva
        self.prs.save(output_path)

        print(f"✅ Apresentação criada: {os.path.basename(output_path)}\n")

        return output_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python converter_profissional.py <arquivo.pdf>")
        sys.exit(1)

    pdf_path = sys.argv[1]
    conversor = ConversorPDFProfissional(pdf_path)
    conversor.converter()
